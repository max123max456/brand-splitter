#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
品牌数据拆分工具 - 云端网页版
用于部署到 Render.com / Railway / Fly.io 等免费云平台
"""
from flask import Flask, request, send_file, render_template_string, jsonify, send_from_directory
import os, pandas as pd, openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from io import BytesIO
import zipfile
import uuid
import shutil

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 500 * 1024 * 1024
UPLOAD_FOLDER = '/tmp/splits'

if not os.path.exists(UPLOAD_FOLDER):
    os.makedirs(UPLOAD_FOLDER)

HF = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
TB = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
CA = Alignment(horizontal='center', vertical='center')
WF = Alignment(wrap_text=True, vertical='top')

def sh(ws, r, c, v):
    cell = ws.cell(r, c, v)
    cell.font = Font(bold=True); cell.fill = HF; cell.border = TB; cell.alignment = CA

def sd(ws, r, c, v, fmt=None, bold=False):
    cell = ws.cell(r, c, v)
    cell.border = TB; cell.alignment = CA
    if fmt: cell.number_format = fmt
    if bold: cell.font = Font(bold=True)

def wr(ws, r, c, vals):
    for i, (v, fmt, b) in enumerate(vals):
        sd(ws, r, c + i, v, fmt, b)

def fv13(a, b, c, d, e, f, g, h, i, j, k, l, m):
    nr14 = j / i if i > 0 else 0
    nr24 = m / l if l > 0 else 0
    return [
        (a, None, False), (b, None, False),
        (int(c), '#,##0', False), (int(d), '#,##0', False),
        (int(e), '#,##0', False), (int(f), '#,##0', False), (int(g), '#,##0', False),
        (h, '#,##0.00', False),
        (int(i), '#,##0', False), (int(j), '#,##0', False),
        (nr14, '0%', False),
        (k, '#,##0.00', False),
        (nr24, '0%', False)
    ]

def fv12(a, b, c, d, e, f, g, h, i, j, k, l):
    nr14 = i / h if h > 0 else 0
    nr24 = l / k if k > 0 else 0
    return [
        (a, None, False),
        (int(b), '#,##0', False), (int(c), '#,##0', False),
        (int(d), '#,##0', False), (int(e), '#,##0', False), (int(f), '#,##0', False),
        (g, '#,##0.00', False),
        (int(h), '#,##0', False), (int(i), '#,##0', False),
        (nr14, '0%', False),
        (j, '#,##0.00', False),
        (nr24, '0%', False)
    ]

def hdr_sum11(ws, r):
    for c, h in enumerate(['品牌', '供应商编号', '提报预算', '品牌UV', '商品详情页UV',
        '14天销售额', '14天ROI', '14天新客率', '24小时销售额', '24HROI', '24H新客率'], 1):
        sh(ws, r, c, h)

def hdr_agg13(ws, r):
    sh(ws, r, 1, '品牌/店铺'); sh(ws, r, 2, '供应商编号')
    for c, h in enumerate(['品牌UV', '商品详情页UV', '14天收藏数', '14天加购数', '14天成交单量',
        '14天销售额', '14天成交客户数', '14天成交新客数', '新客率', '24小时销售额', '24H新客率'], 3):
        sh(ws, r, c, h)

def hdr_agg12(ws, r, fc='名称'):
    sh(ws, r, 1, fc)
    for c, h in enumerate(['品牌UV', '商品详情页UV', '14天收藏数', '14天加购数', '14天成交单量',
        '14天销售额', '14天成交客户数', '14天成交新客数', '新客率', '24小时销售额', '24H新客率'], 2):
        sh(ws, r, c, h)


def run_split(in_file, out_dir, cfg, excl_list):
    def log(msg):
        print(msg)
    try:
        log('[1/5] 读取数据表...')
        sheets = pd.read_excel(in_file, sheet_name=None, engine='openpyxl')
        main_sh = cfg.get('main_sheet', list(sheets.keys())[0])
        df = sheets[main_sh]
        log('  已加载: ' + main_sh + ' (' + str(len(df)) + '行)')

        bc  = cfg['brand_col'];    sc  = cfg['supplier_col']
        dc  = cfg['date_col'];     rc  = cfg['resource_col']
        gc  = cfg['crowd_col'];    uv  = cfg['uv_col']
        uvd = cfg['uvd_col'];      f14 = cfg['fav14_col']
        c14 = cfg['cart14_col'];   o14 = cfg['order14_col']
        sa14= cfg['sales14_col'];  cu14= cfg['cu14_col']
        cu24= cfg.get('cu24_col', '')
        ne14= cfg['ne14_col'];     sa24= cfg['sales24_col']
        ne24= cfg['ne24_col']

        if cu24 and cu24 in df.columns: log('  24H成交客户数列: ' + cu24 + ' [OK]')
        else:
            cu24 = None
            log('  [INFO] 未找到24H成交客户数，将使用14天口径')

        log('[2/5] 读取预算...')
        budget_map = {}
        bs = cfg.get('budget_sheet', '')
        if bs and bs in sheets:
            df_bs = pd.read_excel(in_file, sheet_name=bs, header=None)
            for row in range(4, 51):
                b_raw = df_bs.iloc[row, 0]; s_raw = df_bs.iloc[row, 1]; bud_raw = df_bs.iloc[row, 2]
                if pd.isna(b_raw): continue
                b = str(b_raw).strip()
                if b in ('品牌', '总计', '品牌/店铺', ''): continue
                try: sid = str(int(float(s_raw)))
                except: sid = ''
                try: bgt = float(bud_raw)
                except: bgt = 0
                if b in budget_map:
                    budget_map[b]['budget'] += bgt
                    if not budget_map[b]['sid'] and sid: budget_map[b]['sid'] = sid
                else:
                    budget_map[b] = {'sid': sid, 'budget': bgt}
            log('  预算条目: ' + str(len(budget_map)) + '条')
        else:
            log('  预算表未配置，跳过')

        brands = [b for b in df[bc].unique()
                 if pd.notna(b) and b not in excl_list and str(b) != '全部']
        log('[3/5] 品牌数量: ' + str(len(brands)))
        os.makedirs(out_dir, exist_ok=True)

        log('[4/5] 开始拆分...')
        for idx, brand in enumerate(brands, 1):
            log('  [' + str(idx) + '/' + str(len(brands)) + '] ' + str(brand)[:30])

            db = df[df[bc] == brand].copy()
            t_uv   = int(db[uv].sum())   if uv   in db.columns else 0
            t_uvd  = int(db[uvd].sum())  if uvd  in db.columns else 0
            t_f14  = int(db[f14].sum())  if f14  in db.columns else 0
            t_c14  = int(db[c14].sum())  if c14  in db.columns else 0
            t_o14  = int(db[o14].sum())  if o14  in db.columns else 0
            t_s14  = db[sa14].sum()      if sa14 in db.columns else 0
            t_s24  = db[sa24].sum()      if sa24 in db.columns else 0
            t_cu14 = int(db[cu14].sum()) if cu14 in db.columns else 0
            t_ne14 = int(db[ne14].sum()) if ne14 in db.columns else 0
            t_ne24 = int(db[ne24].sum()) if ne24 in db.columns else 0
            t_cu24 = int(db[cu24].sum()) if cu24 and cu24 in db.columns else t_cu14

            bi  = budget_map.get(brand, {'sid': '', 'budget': 0})
            bud = bi['budget']
            roi14 = t_s14 / bud if bud > 0 else 0
            roi24 = t_s24 / bud if bud > 0 else 0
            nr14  = t_ne14 / t_cu14 if t_cu14 > 0 else 0
            nr24  = t_ne24 / t_cu24 if t_cu24 > 0 else 0

            ag_cols = [uv, uvd, f14, c14, o14, sa14, cu14, ne14, sa24, ne24]
            if cu24: ag_cols.append(cu24)
            ag = {k: 'sum' for k in ag_cols if k in db.columns}
            daily = db.groupby(dc).agg(ag).reset_index() if dc in db.columns and len(db) > 0 else pd.DataFrame()
            res   = db.groupby(rc).agg(ag).reset_index() if rc in db.columns and len(db) > 0 else pd.DataFrame()
            crd   = db.groupby(gc).agg(ag).reset_index() if gc in db.columns and len(db) > 0 else pd.DataFrame()
            if sa14 in crd.columns and len(crd) > 0:
                crd = crd.sort_values(sa14, ascending=False)

            bd = daily.loc[daily[uv].idxmax()] if len(daily) > 0 and uv in daily.columns else None
            br = res.loc[res[sa14].idxmax()]  if len(res) > 0 and sa14 in res.columns else None
            try:    bd_str = bd[dc].strftime('%m/%d') if bd is not None and hasattr(bd[dc], 'strftime') else str(bd[dc])[:10] if bd is not None else ''
            except: bd_str = str(bd[dc])[:10] if bd is not None else ''
            br_str = str(br[rc]) if br is not None else ''
            bc_str = str(crd.iloc[0][gc]) if len(crd) > 0 else ''

            insight = (
                '品牌投放数据分析\n'
                '整体预算' + str(int(bud)) + '，品牌UV' + str(t_uv) + '，14天ROI' + ('%.2f' % roi14) + '，新客率' + str(int(nr14 * 100)) + '%\n'
                '1、分日：' + bd_str + ' 整体UV最高，建议品牌集中投放活动当日，适当蓄水，提升活动曝光引流\n'
                '2、资源位：' + br_str + ' 销售产出最高，适合单品打爆销售提升，搭配高引流资源位，整体拉升活动销售\n'
                '3、人群：' + bc_str + ' 销售相对较高，建议品牌前期重点积累流量，提升品牌UV，逐步拉升转化'
            )

            wb = openpyxl.Workbook(); ws = wb.active; row = 1
            ws.cell(row, 1, insight); ws.merge_cells('A1:M1')
            ws.cell(row, 1).alignment = WF; ws.row_dimensions[row].height = 65; row = 3

            hdr_sum11(ws, row); row += 1
            sd(ws, row, 1, brand); sd(ws, row, 2, bi['sid'])
            sd(ws, row, 3, bud, '#,##0'); sd(ws, row, 4, t_uv, '#,##0'); sd(ws, row, 5, t_uvd, '#,##0')
            sd(ws, row, 6, t_s14, '#,##0.00'); sd(ws, row, 7, roi14, '0.00')
            sd(ws, row, 8, nr14, '0%')
            sd(ws, row, 9, t_s24, '#,##0.00'); sd(ws, row, 10, roi24, '0.00')
            sd(ws, row, 11, nr24, '0%')
            row += 3

            hdr_agg13(ws, row); row += 1
            sup = db.groupby(sc).agg(ag).reset_index() if sc in db.columns and len(db) > 0 else pd.DataFrame()
            for _, sr in sup.iterrows():
                cu14s = int(sr.get(cu14, 0))
                cu24s = int(sr.get(cu24, cu14s)) if cu24 and cu24 in sr.index else cu14s
                ne14s = int(sr.get(ne14, 0)); ne24s = int(sr.get(ne24, 0))
                rv = fv13(brand, str(sr[sc]),
                    sr.get(uv,0), sr.get(uvd,0), sr.get(f14,0), sr.get(c14,0), sr.get(o14,0),
                    sr.get(sa14,0), cu14s, ne14s,
                    sr.get(sa24,0), cu24s, ne24s)
                wr(ws, row, 1, rv); row += 1
            rv = fv13('总计', ' ',
                t_uv, t_uvd, t_f14, t_c14, t_o14, t_s14, t_cu14, t_ne14, t_s24, t_cu24, t_ne24)
            wr(ws, row, 1, rv); row += 2

            ws.cell(row, 1, '品牌/店铺'); ws.cell(row, 2, brand); row += 2
            hdr_agg12(ws, row, '日期'); row += 1
            for _, dr in daily.iterrows():
                try: ds = dr[dc].strftime('%Y-%m-%d')
                except: ds = str(dr[dc])[:10]
                cu14d = int(dr.get(cu14, 0))
                cu24d = int(dr.get(cu24, cu14d)) if cu24 and cu24 in dr.index else cu14d
                ne14d = int(dr.get(ne14, 0)); ne24d = int(dr.get(ne24, 0))
                rv = fv12(ds,
                    dr.get(uv,0), dr.get(uvd,0), dr.get(f14,0), dr.get(c14,0), dr.get(o14,0),
                    dr.get(sa14,0), cu14d, ne14d,
                    dr.get(sa24,0), cu24d, ne24d)
                wr(ws, row, 1, rv); row += 1
            rv = fv12('总计',
                t_uv, t_uvd, t_f14, t_c14, t_o14, t_s14, t_cu14, t_ne14, t_s24, t_cu24, t_ne24)
            wr(ws, row, 1, rv); row += 2

            ws.cell(row, 1, '品牌/店铺'); ws.cell(row, 2, brand); row += 2
            hdr_agg12(ws, row, '投放资源'); row += 1
            for _, rr in res.iterrows():
                cu14r = int(rr.get(cu14, 0))
                cu24r = int(rr.get(cu24, cu14r)) if cu24 and cu24 in rr.index else cu14r
                ne14r = int(rr.get(ne14, 0)); ne24r = int(rr.get(ne24, 0))
                rv = fv12(str(rr[rc]),
                    rr.get(uv,0), rr.get(uvd,0), rr.get(f14,0), rr.get(c14,0), rr.get(o14,0),
                    rr.get(sa14,0), cu14r, ne14r,
                    rr.get(sa24,0), cu24r, ne24r)
                wr(ws, row, 1, rv); row += 1
            rv = fv12('总计',
                t_uv, t_uvd, t_f14, t_c14, t_o14, t_s14, t_cu14, t_ne14, t_s24, t_cu24, t_ne24)
            wr(ws, row, 1, rv); row += 2

            ws.cell(row, 1, '品牌/店铺'); ws.cell(row, 2, brand); row += 2
            hdr_agg12(ws, row, '人群'); row += 1
            for _, cr in crd.iterrows():
                cu14c = int(cr.get(cu14, 0))
                cu24c = int(cr.get(cu24, cu14c)) if cu24 and cu24 in cr.index else cu14c
                ne14c = int(cr.get(ne14, 0)); ne24c = int(cr.get(ne24, 0))
                rv = fv12(str(cr[gc]),
                    cr.get(uv,0), cr.get(uvd,0), cr.get(f14,0), cr.get(c14,0), cr.get(o14,0),
                    cr.get(sa14,0), cu14c, ne14c,
                    cr.get(sa24,0), cu24c, ne24c)
                wr(ws, row, 1, rv); row += 1
            rv = fv12('总计',
                t_uv, t_uvd, t_f14, t_c14, t_o14, t_s14, t_cu14, t_ne14, t_s24, t_cu24, t_ne24)
            wr(ws, row, 1, rv)

            for c in range(1, 14): ws.column_dimensions[get_column_letter(c)].width = 15
            ws.column_dimensions['B'].width = 12

            safe = ''.join(ch for ch in str(brand) if ch not in '\\/:*?"<>|')
            out_path = os.path.join(out_dir, safe + '.xlsx')
            wb.save(out_path)
            log('    >> ' + safe + '.xlsx')

        log('[5/5] 完成！共生成 ' + str(len(brands)) + ' 个文件')
        return True, '完成！共生成 ' + str(len(brands)) + ' 个文件'
    except Exception as e:
        import traceback; traceback.print_exc()
        return False, '错误: ' + str(e)


HTML = '''
<!DOCTYPE html>
<html lang="zh">
<head>
<meta charset="utf-8">
<title>品牌数据拆分工具</title>
<style>
*{box-sizing:border-box;margin:0;padding:0}
body{font-family:-apple-system,BlinkMacSystemFont,'Segoe UI',Roboto,sans-serif;background:#f0f2f5;min-height:100vh}
header{background:linear-gradient(135deg,#667eea,#764ba2);color:#fff;padding:24px 30px}
header h1{font-size:22px;font-weight:700}
header p{font-size:14px;opacity:0.9;margin-top:4px}
.container{max-width:800px;margin:20px auto;padding:0 16px}
.card{background:#fff;border-radius:12px;box-shadow:0 2px 8px rgba(0,0,0,0.08);padding:24px;margin-bottom:16px}
.card h2{font-size:15px;color:#333;margin-bottom:14px;font-weight:600}
.form-group{margin-bottom:14px}
.form-group label{display:block;font-size:13px;color:#555;margin-bottom:6px}
.form-group input{width:100%;padding:10px 12px;border:1px solid #e0e0e0;border-radius:8px;font-size:14px;transition:border-color 0.2s}
.form-group input:focus{outline:none;border-color:#667eea}
.form-row{display:flex;gap:10px}
.form-row .form-group{flex:1}
.file-drop{border:2px dashed #ddd;border-radius:10px;padding:36px 20px;text-align:center;cursor:pointer;transition:all 0.2s;background:#fafafa}
.file-drop:hover{border-color:#667eea;background:#f8f6ff}
.file-drop.dragover{border-color:#667eea;background:#f0edff}
.file-drop input{display:none}
.file-icon{font-size:40px}
.file-text{font-size:14px;color:#666;margin-top:8px}
.file-name{font-size:13px;color:#333;margin-top:8px;font-weight:500}
.btn{display:inline-flex;align-items:center;justify-content:center;gap:6px;padding:12px 24px;border:none;border-radius:8px;font-size:14px;font-weight:500;cursor:pointer;transition:all 0.2s}
.btn-primary{background:linear-gradient(135deg,#667eea,#764ba2);color:#fff}
.btn-primary:hover{transform:translateY(-1px);box-shadow:0 4px 12px rgba(102,126,234,0.4)}
.btn-primary:disabled{opacity:0.6;cursor:not-allowed;transform:none;box-shadow:none}
.btn-secondary{background:#f1f3f4;color:#555;border:1px solid #e0e0e0}
.btn-group{display:flex;gap:10px;margin-top:16px}
.log-box{background:#1e1e1e;color:#d4d4d4;padding:16px;border-radius:8px;font-size:12px;line-height:1.7;max-height:300px;overflow-y:auto;font-family:'SF Mono',Monaco,monospace;margin-top:16px}
.progress{margin-top:16px;display:none}
.progress-bar{height:4px;background:#e0e0e0;border-radius:2px;overflow:hidden}
.progress-fill{height:100%;background:linear-gradient(90deg,#667eea,#764ba2);width:0;transition:width 0.3s}
.status{font-size:13px;color:#666;margin-top:8px;text-align:center}
.result{background:#e8f5e9;border-radius:8px;padding:16px;margin-top:16px;display:none}
.result.show{display:block}
.result a{color:#2e7d32;font-weight:500}
.note{font-size:12px;color:#888;margin-top:12px;line-height:1.6}
</style>
</head>
<body>
<header>
  <h1>品牌数据拆分工具</h1>
  <p>在线版 · Mac/Windows/Android/iOS 通用 · 无需安装任何软件</p>
</header>
<div class="container">
  <div class="card">
    <h2>📁 选择Excel文件</h2>
    <div class="file-drop" id="dropZone" onclick="document.getElementById('fileInput').click()">
      <input type="file" id="fileInput" accept=".xlsx,.xls">
      <div class="file-icon">☁️</div>
      <div class="file-text">��击选择或将文件拖放到这里</div>
      <div class="file-name" id="fileName"></div>
    </div>
  </div>

  <div class="card">
    <h2>⚙️ 列名配置（一般不用改）</h2>
    <div class="form-row">
      <div class="form-group"><label>主数据Sheet</label><input id="main_sheet" value="投放数据源"></div>
      <div class="form-group"><label>预算Sheet</label><input id="budget_sheet" value="整体数据"></div>
    </div>
    <div class="form-row">
      <div class="form-group"><label>品牌列</label><input id="brand_col" value="品牌/店铺"></div>
      <div class="form-group"><label>供应商列</label><input id="supplier_col" value="供应商编号"></div>
    </div>
    <div class="form-group"><label>排除品牌</label><input id="exclude_brands" value="全部" placeholder="多个用逗号分隔"></div>
    <div class="note">💡 如列名不对，点击"检测列名"自动识别</div>
  </div>

  <div class="btn-group">
    <button class="btn btn-primary" id="btnStart" onclick="startSplit()" disabled>▶ 开始拆分</button>
    <button class="btn btn-secondary" onclick="detectCols()">🔍 检测列名</button>
  </div>

  <div class="progress" id="progress">
    <div class="progress-bar"><div class="progress-fill" id="progressFill"></div></div>
    <div class="status" id="status">准备中...</div>
  </div>

  <div class="log-box" id="logBox"></div>

  <div class="result" id="result">
    <p>✅ 拆分完成！共生成 <span id="count">0</span> 个文件</p>
    <p style="margin-top:8px"><a id="downloadLink" href="#">📥 点击下载结果</a></p>
  </div>
</div>

<script>
let selectedFile = null;
const logBox = document.getElementById('logBox');
const btn = document.getElementById('btnStart');
const progress = document.getElementById('progress');
const progressFill = document.getElementById('progressFill');
const status = document.getElementById('status');

function log(msg, type='info') {
  const cls = type === 'error' ? 'color:#f66' : type === 'ok' ? 'color:#6c6' : '';
  const div = document.createElement('div');
  div.style.cssText = 'margin:2px 0;' + cls;
  div.textContent = new Date().toLocaleTimeString() + ' ' + msg;
  logBox.appendChild(div);
  logBox.scrollTop = logBox.scrollHeight;
}

document.getElementById('fileInput').addEventListener('change', function(e) {
  if (e.target.files.length > 0) {
    selectedFile = e.target.files[0];
    document.getElementById('fileName').textContent = selectedFile.name + ' (' + (selectedFile.size/1024/1024).toFixed(1) + ' MB)';
    btn.disabled = false;
    log('已选择: ' + selectedFile.name);
  }
});

const dropZone = document.getElementById('dropZone');
dropZone.addEventListener('dragover', function(e){e.preventDefault();dropZone.classList.add('dragover')});
dropZone.addEventListener('dragleave', function(){dropZone.classList.remove('dragover')});
dropZone.addEventListener('drop', function(e){
  e.preventDefault();dropZone.classList.remove('dragover');
  if(e.dataTransfer.files.length>0){
    selectedFile=e.dataTransfer.files[0];
    document.getElementById('fileName').textContent=selectedFile.name+' ('+(selectedFile.size/1024/1024).toFixed(1)+' MB)';btn.disabled=false;log('已选择: '+selectedFile.name)}});

function detectCols(){
  if(!selectedFile){alert('请先选择文件');return}
  log('检测列名...');
  const fd=new FormData();fd.append('file',selectedFile);
  fetch('/detect',{method:'POST',body:fd}).then(r=>r.json()).then(d=>{
    if(d.error){log(d.error,'error');return}
    for(const[k,v] of Object.entries(d.cols)){
      const el=document.getElementById(k);if(el)el.value=v}
    log('列名: '+Object.keys(d.cols).length+'个','ok')})}

function startSplit(){
  if(!selectedFile){alert('请先选择文件');return}
  logBox.innerHTML='';btn.disabled=true;progress.style.display='block';progressFill.style.width='0%';status.textContent='上传中...';
  log('开始拆分...');
  const fd=new FormData();fd.append('file',selectedFile);
  fd.append('cfg_main_sheet',document.getElementById('main_sheet').value);
  fd.append('cfg_budget_sheet',document.getElementById('budget_sheet').value);
  fd.append('cfg_brand_col',document.getElementById('brand_col').value);
  fd.append('cfg_supplier_col',document.getElementById('supplier_col').value);
  fd.append('cfg_exclude_brands',document.getElementById('exclude_brands').value);

  fetch('/split',{method:'POST',body:fd}).then(r=>r.json()).then(d=>{
    btn.disabled=false;
    if(d.success){
      log(d.message,'ok');
      document.getElementById('count').textContent=d.count;
      document.getElementById('downloadLink').href=d.download_url;
      document.getElementById('result').classList.add('show')}
    else{log(d.message,'error')}
  }).catch(e=>{btn.disabled=false;log('错误: '+e,'error')})}
</script>
</body>
</html>
'''


@app.route('/')
def index():
    return render_template_string(HTML)


@app.route('/detect', methods=['POST'])
def detect():
    if 'file' not in request.files:
        return jsonify({'error': '未上传文件'})
    f = request.files['file']
    try:
        sheets = pd.read_excel(f, sheet_name=None, engine='openpyxl')
        main_sh = list(sheets.keys())[0]
        return jsonify({'sheets': list(sheets.keys()), 'cols': {c: c for c in sheets[main_sh].columns}})
    except Exception as e:
        return jsonify({'error': str(e)})


@app.route('/split', methods=['POST'])
def split():
    if 'file' not in request.files:
        return jsonify({'success': False, 'message': '未上传文件'})
    f = request.files['file']
    
    # 生成唯一ID
    job_id = str(uuid.uuid4())[:8]
    tmp_dir = os.path.join(UPLOAD_FOLDER, job_id)
    os.makedirs(tmp_dir, exist_ok=True)
    
    # 保存上传文件
    in_file = os.path.join(tmp_dir, 'input.xlsx')
    f.save(in_file)
    
    cfg = {
        'main_sheet': request.form.get('cfg_main_sheet', '投放数据源'),
        'budget_sheet': request.form.get('cfg_budget_sheet', '整体数据'),
        'brand_col': request.form.get('cfg_brand_col', '品牌/店铺'),
        'supplier_col': request.form.get('cfg_supplier_col', '供应商编号'),
    }
    
    # 添加默认列名
    for col in ['date_col', 'resource_col', 'crowd_col', 'uv_col', 'uvd_col', 'fav14_col', 'cart14_col', 
              'order14_col', 'sales14_col', 'cu14_col', 'ne14_col', 'sales24_col', 'ne24_col']:
        cfg[col] = col.replace('_col', '')
        if '24' in col:
            cfg[col] = col.replace('_col', '').replace('14', '24')
        else:
            cfg[col] = cfg.get(col, col.replace('_col', ''))
    
    out_dir = os.path.join(tmp_dir, 'output')
    
    try:
        ok, msg = run_split(in_file, out_dir, cfg, ['全部'])
        
        if ok:
            # 创建ZIP
            zip_name = f'品牌数据_{job_id}.zip'
            zip_path = os.path.join(UPLOAD_FOLDER, zip_name)
            with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as zf:
                for root, dirs, files in os.walk(out_dir):
                    for file in files:
                        if file.endswith('.xlsx'):
                            fp = os.path.join(root, file)
                            zf.write(fp, file)
            
            count = len([f for f in os.listdir(out_dir) if f.endswith('.xlsx')])
            
            # 清理临时文件
            os.remove(in_file)
            shutil.rmtree(out_dir)
            
            return jsonify({
                'success': True, 
                'message': msg,
                'count': count,
                'download_url': f'/download/{zip_name}'
            })
        else:
            return jsonify({'success': False, 'message': msg})
            
    except Exception as e:
        import traceback; traceback.print_exc()
        return jsonify({'success': False, 'message': str(e)})


@app.route('/download/<name>')
def download(name):
    return send_from_directory(UPLOAD_FOLDER, name, as_attachment=True)


if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5188))
    app.run(host='0.0.0.0', port=port, debug=False)